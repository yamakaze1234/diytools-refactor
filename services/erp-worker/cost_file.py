"""只读成本文件。只消费明确指定的两列；不启动 Excel、不执行宏或外部刷新。"""
import hashlib
import math
import copy
import time
_CACHE = None
from pathlib import Path
from decimal import Decimal, InvalidOperation

def read_costs(settings):
    global _CACHE
    started=time.perf_counter()
    from pyxlsb import open_workbook
    p=Path(settings.get('path',''))
    if not p.is_absolute() or p.suffix.lower() not in {'.xlsb','.xlsx'}:
        raise ValueError('请选择绝对路径的 xlsb 或 xlsx 表格')
    if p.stat().st_size>100*1024*1024:
        raise ValueError('成本表超过 100 MiB 限制')
    def fingerprint():
        with p.open('rb') as f:return hashlib.file_digest(f,'sha256').hexdigest()
    sha=fingerprint();sheet=settings.get('sheet','');gc=settings.get('goodsColumn',0)-1;tc=settings.get('taxColumn',0)-1
    if not all(type(n) is int and 0<=n<16384 for n in [gc,tc]) or gc==tc:raise ValueError('列号无效')
    # 仅缓存最近一份已验证的两列结果；摘要变化即失效，不保留工作簿对象。
    key=(str(p),sha,sheet,gc,tc)
    if _CACHE is not None and _CACHE[0]==key:
        if fingerprint()!=sha:raise ValueError('读取期间表格发生变化，请重新读取')
        result=copy.deepcopy(_CACHE[1]);result['cacheHit']=True;result['fileReadMs']=round((time.perf_counter()-started)*1000);return result
    _CACHE=None
    values={};skipped=0;invalid=[];header=False;count=0
    def consume(a,b):
        nonlocal skipped,header,count
        count+=1
        if count>200000:raise ValueError('工作表超过 20 万行限制')
        if not header:
            if str(a).strip().lower()=='goods_id' and str(b).strip()=='含税成本':header=True
            elif count>=30:raise ValueError('前 30 行未找到指定列的 goods_id 和含税成本表头')
            return
        if a is None or str(a).strip() in ('','goods_id') or b is None or str(b).strip()=='':skipped+=1;return
        if isinstance(a,bool) or isinstance(b,bool):raise ValueError('商品 ID 或成本不能为布尔值')
        if isinstance(a,(int,float)):
            if not math.isfinite(a) or a!=int(a) or abs(a)>9007199254740991:raise ValueError('商品 ID 数值不精确，请在 Excel 中存为文本')
            a=str(int(a))
        a=str(a).strip()
        if not a.isascii() or not a.isdigit() or len(a)>100:raise ValueError('goods_id 必须为数字或数字文本')
        try:price=Decimal(str(b).strip())
        except InvalidOperation:
            invalid.append({'goods_id':a,'reason':'含税成本无有效缓存值'});skipped+=1;return
        if not price.is_finite() or not 0<=price<=100000000:raise ValueError('含税成本超出范围')
        price=float(price)
        if a in values and values[a]!=price:raise ValueError('同一 goods_id 存在不同成本：'+a)
        values[a]=price
        if len(values)>100000:raise ValueError('商品超过 10 万条限制')
    if p.suffix.lower()=='.xlsb':
        with open_workbook(str(p)) as w:
            with w.get_sheet(sheet) as s:
                for row in s.rows(sparse=True):
                    selected={c.c:c.v for c in row if c.c in (gc,tc)}
                    consume(selected.get(gc),selected.get(tc))
    else:
        from openpyxl import load_workbook
        w=load_workbook(p,read_only=True,data_only=True,keep_links=False)
        try:
            for row in w[sheet].iter_rows(min_col=min(gc,tc)+1,max_col=max(gc,tc)+1,values_only=True):consume(row[gc-min(gc,tc)],row[tc-min(gc,tc)])
        finally:w.close()
    if not header or not values:raise ValueError('未读取到有效成本，请核对工作表、列号和已保存的公式结果')
    if fingerprint()!=sha:raise ValueError('读取期间表格发生变化，请重新读取')
    result={'rows':[{'goods_id':i,'tax':v} for i,v in values.items()],'sourceSha256':sha,'skipped':skipped,'total':len(values),'invalid':invalid[:80],'invalidCount':len(invalid)}

    _CACHE=(key,copy.deepcopy(result));result['cacheHit']=False;result['fileReadMs']=round((time.perf_counter()-started)*1000);return result
